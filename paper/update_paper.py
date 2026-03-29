"""논문 업데이트 스크립트 — 매칭 검증, Nested Logit, 이미지 추가."""
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pathlib import Path
from copy import deepcopy

PAPER_DIR = Path(__file__).resolve().parent
ROOT = PAPER_DIR.parent
IMG_DIR = ROOT / 'images'

doc = Document(str(PAPER_DIR / 'paper_draft.docx'))


def find_para_index(doc, text_contains):
    for i, p in enumerate(doc.paragraphs):
        if text_contains in p.text:
            return i
    return None


def add_paragraph_after(doc, after_index, text, style_name=None):
    """after_index 단락 뒤에 새 단락 삽입. 반환: 새 단락의 인덱스."""
    ref_para = doc.paragraphs[after_index]
    new_para = doc.add_paragraph(text)
    if style_name:
        new_para.style = doc.styles[style_name]
    # XML 수준에서 ref 바로 뒤로 이동
    ref_para._element.addnext(new_para._element)
    return after_index + 1


def add_image_after(doc, after_index, img_path, width_inches=6.0, caption=''):
    """after_index 단락 뒤에 이미지 + 캡션 삽입."""
    ref_para = doc.paragraphs[after_index]

    # 이미지 단락
    img_para = doc.add_paragraph()
    img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = img_para.add_run()
    run.add_picture(str(img_path), width=Inches(width_inches))
    ref_para._element.addnext(img_para._element)
    after_index += 1

    # 캡션
    if caption:
        cap_para = doc.add_paragraph(caption)
        cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cap_para.runs:
            r.font.size = Pt(9)
        img_para._element.addnext(cap_para._element)
        after_index += 1

    return after_index


# ============================================================
# 1. Fig.2 시퀀스 유사도 예시 (3.4.1 뒤, 3.4.2 앞)
# ============================================================
idx = find_para_index(doc, '3.4.2 매칭 임계값')
if idx:
    img_path = IMG_DIR / 'fig2_sequence_similarity_examples.png'
    if img_path.exists():
        add_image_after(doc, idx - 1, img_path, 6.0,
            'Fig. 2. 유사도 수준별 AFC-OTP 경로 매칭 예시. (a) LCS=1.00: 완전 일치, (b) LCS=0.67: 주요 노선 일치, (c) LCS=0.50: 중간 경로 상이, (d) LCS=0.33: 수단 조합 상이.')
        print('OK: Fig.2 삽입')

# ============================================================
# 2. Fig.3 히트맵 (3.4.3 뒤, 3.4.4 앞)
# ============================================================
idx = find_para_index(doc, '3.4.4 훈련 데이터 구성')
if idx:
    img_path = IMG_DIR / 'fig3_2d_grid_heatmap.png'
    if img_path.exists():
        add_image_after(doc, idx - 1, img_path, 6.0,
            'Fig. 3. 유사도 가중치 2단계 격자탐색 결과. (a) Test McFadden ρ², (b) First Preference Recovery.')
        print('OK: Fig.3 삽입')

# ============================================================
# 3. 매칭 정확도 검증 (4.3 성능 평가 뒤, 4.4 수단분담률 앞)
# ============================================================
idx = find_para_index(doc, '4.4 수단분담률 재현')
if idx:
    cur = idx - 1

    cur = add_paragraph_after(doc, cur, '4.3.1 매칭 정확도 검증', 'Heading 3')
    cur = add_paragraph_after(doc, cur,
        '유사도 매칭의 신뢰성을 검증하기 위해, 전체 526,968개 OD 중 7.2%에 해당하는 37,997개 OD를 '
        '유사도 구간별 층화 추출하여 AFC 통행과 OTP 매칭 경로의 노선·수단·환승횟수를 비교하였다. '
        '자동 판정 기준은 다음과 같다: EXACT(노선·수단·환승 모두 일치), PARTIAL(노선 70% 이상 겹침 또는 수단 일치), '
        'WEAK(노선 30% 이상 겹침), MISMATCH(나머지).')
    cur = add_paragraph_after(doc, cur,
        '검증 결과, 정확(EXACT) 48.7%, 부분일치(PARTIAL) 48.2%로 전체 정확 및 부분일치율이 '
        '96.9%(36,832/37,997)에 달하였다. 불일치(MISMATCH)는 992건(2.6%)으로, 모두 스마트카드에서는 '
        '철도만 이용하였으나 OTP가 접근 버스를 추가 제안한 경우였다. 즉, 주 이용 노선(철도)은 일치하나 접근 수단에서 '
        '차이가 발생한 것으로, 실질적 경로 오류는 아니다.')
    cur = add_paragraph_after(doc, cur,
        '유사도 구간별로는 composite similarity 0.6 이상에서 정확률이 100%였으며, 0.5~0.6 구간에서도 '
        '72.0%의 정확률을 보였다. 개별 항목별 일치율은 수단 일치 96.9%, 환승횟수 일치 90.1%, 노선 완전 일치 '
        '49.7%로 나타났다. 노선명의 정규화 차이(예: OTP "서울2호선" vs SC "2호선")에 의한 부분불일치가 '
        '존재하나, 수단 수준에서는 높은 정확도를 확인하였다.')

    # 매칭 검증 이미지
    img_path = IMG_DIR / 'fig_matching_validation.png'
    if img_path.exists():
        cur = add_image_after(doc, cur, img_path, 6.0,
            'Fig. 4. 매칭 정확도 검증 결과 (n=37,997). (a) 판정 분포, (b) 유사도 vs 노선 겹침 비율, (c) 유사도 구간별 정확률.')
    print('OK: 매칭 검증 섹션 삽입')

# ============================================================
# 4. Nested Logit IIA 검증 (4.5 모형 사양 진화 뒤, 4.6 딥러닝 앞)
# ============================================================
idx = find_para_index(doc, '4.6 딥러닝 모형 비교')
if idx:
    cur = idx - 1

    cur = add_paragraph_after(doc, cur, '4.5.1 IIA 가정 검증 (Nested Logit)', 'Heading 3')
    cur = add_paragraph_after(doc, cur,
        'MNL의 IIA(Independence of Irrelevant Alternatives) 가정의 적절성을 검증하기 위해 Nested Logit '
        '모형을 추정하였다. Nested Logit은 유사한 대안을 nest로 묶어 nest 내 대안 간 상관을 허용하는 모형으로, '
        'nest scale parameter μ가 1이면 MNL과 동일하고, μ < 1이면 IIA 가정이 위반됨을 의미한다.')
    cur = add_paragraph_after(doc, cur,
        '본 연구에서는 주 교통수단을 기준으로 bus, train, gtx의 3개 nest를 구성하였다. 버스+철도 혼합 경로는 '
        'IVT가 더 큰 수단의 nest에 할당하였다. nest 분포는 bus 72.4%, train 26.6%, gtx 1.0%이다.')
    cur = add_paragraph_after(doc, cur,
        '추정 결과, μ_bus = 0.719(t = 300.1)로 1과 유의하게 달랐으며, μ_train = 1.000, μ_gtx = 1.000으로 '
        'MNL과 동일하였다. 이는 버스 nest 내 대안 간에만 약한 상관(IIA 위반)이 존재하고, 철도·GTX nest에서는 '
        'MNL의 독립성 가정이 성립함을 의미한다.')
    cur = add_paragraph_after(doc, cur,
        'Nested Logit의 test ρ² = 0.533, Top-1 FPR 72.4%로 MNL(ρ² = 0.530, FPR 72.5%)과 거의 동일한 '
        '성능을 보였다. 모든 β의 부호가 이론적 기대와 일치하였으며, 파라미터 수가 9개에서 12개로 증가한 데 비해 '
        '성능 개선이 미미하여, 본 데이터에서 MNL의 IIA 가정이 대체로 적절함을 뒷받침한다.')
    print('OK: Nested Logit 섹션 삽입')

# ============================================================
# 5. SHAP 비교 이미지 (4.7.3 뒤)
# ============================================================
idx = find_para_index(doc, '4.7.3 MNL과 딥러닝의 해석 일관성')
if idx:
    # 4.7.3 내용의 마지막 문단 찾기
    last_idx = idx
    for i in range(idx + 1, len(doc.paragraphs)):
        if doc.paragraphs[i].style.name.startswith('Heading'):
            break
        if doc.paragraphs[i].text.strip():
            last_idx = i

    img_path = ROOT / 'data' / 'training_set' / 'shap_comparison.png'
    if img_path.exists():
        add_image_after(doc, last_idx, img_path, 6.0,
            'Fig. 5. SHAP 기반 변수 중요도 비교 (Mean |SHAP value|). 5개 딥러닝 모형에서 공통적으로 수단 종류(Has bus)와 접근 도보시간이 가장 높은 중요도를 보였다.')
        print('OK: SHAP 비교 이미지 삽입')

# ============================================================
# 6. 선행연구 비교 설명 (2.1 마지막, 2.2 앞)
# ============================================================
idx = find_para_index(doc, '2.2 스마트카드 데이터 활용 연구')
if idx:
    add_paragraph_after(doc, idx - 1,
        'Table 1은 기존 대중교통 경로선택 연구와 본 연구의 방법론적 차이를 비교한 것이다. '
        '기존 연구는 GPS 추적이나 SP 조사에 의존하여 표본 규모가 수천~수만 OD로 제한적이었으며, '
        '단일 수단(도시철도) 또는 단일 모형(MNL/PSL)에 집중하는 경향이 있었다. '
        'Marra and Corman(2020)은 취리히에서 GPS 추적 기반 1,388통행으로 choice set 구성 기준을 제안하였고, '
        'Tomhave and Khani(2022)는 미니애폴리스에서 RP 데이터 기반 다기준 경로선택을 분석하였다. '
        'Zhao et al.(2017)은 심천 지하철에서 AFC 기반 확률적 배정을 수행하였으나 도시철도 단독 네트워크에 한정되었다. '
        'Arriagada et al.(2025)는 산티아고에서 스마트카드 기반 경험학습 모형을 제안하였으나 신규 노선 학습에 초점을 맞추었다. '
        'Marra and Corman(2025)은 취리히에서 딥러닝(DNN)과 PSL을 비교하였으나 단일 DL 모형만 분석하였다. '
        '본 연구는 스마트카드 빅데이터(53만 OD)와 RAPTOR 기반 경로탐색엔진을 결합하여 '
        '버스·도시철도·GTX를 포괄하는 다수단 경로선택 모형을 구축하고, '
        'MNL부터 딥러닝 5종(DNN, TasteNet, ResLogit, ASU-DNN, L-MNL)까지 체계적으로 비교한 점에서 차별화된다.')
    print('OK: 선행연구 비교 설명 삽입')

# ============================================================
# 7. 5.5 한계에 Nested Logit 추가
# ============================================================
idx = find_para_index(doc, '둘째, 대기시간 파라미터가 구조적으로 식별되지 않았다')
if idx:
    add_paragraph_after(doc, idx - 1,
        '또한, Nested Logit 검증 결과 μ_bus = 0.719로 버스 대안 간 약한 상관이 확인되었다. '
        '향후 Mixed Logit이나 Cross-Nested Logit을 통해 보다 유연한 대안 간 상관 구조를 반영할 수 있을 것이다.')
    print('OK: 한계 Nested 추가')

# ============================================================
# 8. 5.1 결론에 매칭 검증 + Nested Logit 결과 추가
# ============================================================
idx = find_para_index(doc, '첫째, 노선 Jaccard(0.40)')
if idx:
    add_paragraph_after(doc, idx,
        '셋째, 37,997개 OD의 층화 검증에서 매칭 정확 및 부분일치율 96.9%를 달성하여 유사도 매칭의 신뢰성을 실증하였다. '
        '넷째, Nested Logit 추정 결과 μ_bus = 0.719만 1과 유의하게 달라 MNL의 IIA 가정이 대체로 적절함을 확인하였다.')
    print('OK: 결론 매칭+Nested 추가')

# ============================================================
# 저장
# ============================================================
output_path = PAPER_DIR / 'paper_draft_updated.docx'
doc.save(str(output_path))
print(f'\n저장 완료: {output_path}')
